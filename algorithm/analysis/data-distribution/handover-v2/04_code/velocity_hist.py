import scipy.io as sio, pandas as pd, numpy as np, os, time, shutil
EVR=os.environ["EVEYE_ROOT"]
FE=os.path.join(EVR,"processed_data","Frame_event_pupil_track_result")
DD="/sessions/busy-happy-newton/mnt/DATASET/codes/data-distribution"
SESS={101:"session_1_0_1",102:"session_1_0_2",201:"session_2_0_1",202:"session_2_0_2"}
VSACC=493.0; EWIN=5000.0
EDGES=np.array([0,1,2,3,4,5,6,7,8,9,10, 15,20,25,30,35,40,45,50, 100,200,300,400,800,1200, np.inf])
LAB=["0-1","1-2","2-3","3-4","4-5","5-6","6-7","7-8","8-9","9-10",
     "10-15","15-20","20-25","25-30","30-35","35-40","40-45","45-50",
     "50-100","100-200","200-300","300-400","400-800","800-1200","1200+"]
NB=len(LAB)  # 25
v=pd.read_csv(DD+"/cache/vel4state_sessions.csv")[["user","eye","code"]].drop_duplicates()

def vels(f):
    M=np.array(sio.loadmat(f)["matcell"],dtype=float)
    fa=M[M[:,5]==1]; fa=fa[np.argsort(fa[:,2])]          # FRAME: 40ms anchors, literal
    tf=fa[:,2]/1e6; dtf=np.diff(tf)
    vf=np.hypot(np.diff(fa[:,3]),np.diff(fa[:,4]))/dtf; vf=vf[np.isfinite(vf)&(dtf>0)]
    Ms=M[np.argsort(M[:,2])]; t=Ms[:,2]                   # EVENT: 5ms grid I-VT, median-9 smoothed
    xs=pd.Series(Ms[:,3]).rolling(9,center=True,min_periods=1).median().values
    ys=pd.Series(Ms[:,4]).rolling(9,center=True,min_periods=1).median().values
    g=np.arange(t.min(),t.max(),EWIN)
    if len(g)>=3:
        j=np.clip(np.searchsorted(t,g),0,len(t)-1)
        ve=np.hypot(np.diff(xs[j]),np.diff(ys[j]))/(EWIN/1e6); ve=ve[np.isfinite(ve)]
    else: ve=np.array([])
    return vf,ve

acc={u:{k:np.zeros(NB) for k in ["frraw","frw","evraw","evw"]} for u in range(1,49)}
t0=time.time()
for _,r in v.iterrows():
    f=os.path.join(FE,r.eye,f"update_20_point_user{int(r.user)}_{SESS[int(r.code)]}.mat")
    if not os.path.exists(f): continue
    vf,ve=vels(f); u=int(r.user)
    acc[u]["frraw"]+=np.histogram(vf,EDGES)[0]
    acc[u]["frw"] +=np.histogram(vf,EDGES,weights=np.where(vf>VSACC,8,2))[0]
    acc[u]["evraw"]+=np.histogram(ve,EDGES)[0]
    acc[u]["evw"] +=np.histogram(ve,EDGES,weights=np.where(ve>VSACC,8,2))[0]
print("pass %.1fs"%(time.time()-t0))

RLAB=[l+" %" for l in LAB]
def df_of(key):
    rows=[]
    for u in range(1,49):
        a=acc[u][key].astype(float); tot=a.sum()
        rec={"Subject":u,"Split":"Train" if u<=36 else "Test","InVal":u in {37,40,45,48}}
        for i in range(NB): rec[LAB[i]]=int(a[i])
        rec["Total"]=int(tot)
        for i in range(NB): rec[RLAB[i]]=round(100*a[i]/tot,3) if tot>0 else 0.0
        rows.append(rec)
    cols=["Subject","Split","InVal"]+LAB+["Total"]+RLAB
    return pd.DataFrame(rows)[cols]
FR,FRi,EV,EVi=df_of("frraw"),df_of("frw"),df_of("evraw"),df_of("evw")

def ssum(df,tag):
    out=[]
    for sp,us in [("Train",set(range(1,37))),("Val",{37,40,45,48}),("Test",set(range(37,49)))]:
        d=df[df.Subject.isin(us)][LAB].sum(); tot=d.sum()
        out.append({"set":tag,"split":sp,"Total":int(tot),**{l:round(100*d[l]/tot,3) for l in LAB}})
    return pd.DataFrame(out)
summary=pd.concat([ssum(FR,"Frame_RAW"),ssum(FRi,"Frame_Interp"),ssum(EV,"Event_RAW"),ssum(EVi,"Event_Interp")],ignore_index=True)
for nm,d in [("frame_raw",FR),("frame_interp",FRi),("event_raw",EV),("event_interp",EVi)]:
    d.to_csv(DD+f"/tables/tbl_vel_{nm}.csv",index=False)
readme=pd.DataFrame({"item":["bins","layout","Frame velocity","Event velocity","Vsacc","RAW","Interpolated(보간)","caveat"],
 "value":["0-10 @1px/s(10) | 10-50 @5px/s(8) | 50-100 | 100-200 | 200-300 | 300-400 | 400-800 | 800-1200 | >=1200  (25 cols)",
          "each sheet: [25 Count cols] + Total + [25 Ratio% cols (suffix ' %')]; ratio = bin/Total*100 per subject",
          "frame centers @25Hz (matcell col5==1, 40ms): v=hypot(dx,dy)/dt, no smoothing",
          "event track 5ms grid (200Hz) I-VT on median-9 smoothed positions: v=hypot(dx,dy)/5ms",
          "493 px/s",
          "per-sample velocity counts (weight 1)",
          "augmentation-weighted: v>493 ->x8, else ->x2 (Blink excluded). same x2/x8 policy; ratio vs weighted Total",
          "native ~0.27ms event diffs jitter-dominated (unusable); 5ms grid used. >=1200 event bin may hold rare jitter outliers"]})
tmp="/tmp/Velocity_dist.xlsx"
with pd.ExcelWriter(tmp,engine="openpyxl") as w:
    readme.to_excel(w,sheet_name="README",index=False)
    FR.to_excel(w,sheet_name="Frame_RAW",index=False); FRi.to_excel(w,sheet_name="Frame_Interp",index=False)
    EV.to_excel(w,sheet_name="Event_RAW",index=False); EVi.to_excel(w,sheet_name="Event_Interp",index=False)
    summary.to_excel(w,sheet_name="summary_split",index=False)
import openpyxl
from openpyxl.styles import Font, PatternFill
cfill=PatternFill("solid",fgColor="E8F0FE"); rfill=PatternFill("solid",fgColor="FCE8E6")
wb=openpyxl.load_workbook(tmp)
for ws in wb.worksheets:
    for cl in ws[1]: cl.font=Font(bold=True)
    if ws.title in ("Frame_RAW","Frame_Interp","Event_RAW","Event_Interp"):
        ws.freeze_panes="D2"
        # tint count headers blue, ratio headers red for readability
        for j,cell in enumerate(ws[1]):
            h=cell.value
            if h in LAB: cell.fill=cfill
            elif isinstance(h,str) and h.endswith(" %"): cell.fill=rfill
    else: ws.freeze_panes="A2"
wb.save(tmp)
shutil.copy(tmp,DD+"/Velocity_distribution.xlsx")
shutil.copy(tmp,"/sessions/busy-happy-newton/mnt/outputs/Velocity_distribution.xlsx")
print("WROTE Velocity_distribution.xlsx  (cols/sheet = 3 + %d count + 1 Total + %d ratio = %d)"%(NB,NB,3+NB+1+NB))
print("\n=== Frame_RAW subj1 counts(0-10 @1px) ===")
print(FR.loc[0,["Subject"]+LAB[:10]].to_string())
print("\n=== summary Train Event_RAW vs Event_Interp (first 12 bins %) ===")
print(summary[(summary.split=='Train')&(summary.set.isin(['Event_RAW','Event_Interp']))].set_index('set')[LAB[:12]].to_string())
